from pathlib import Path
import concurrent.futures
import csv
import hashlib
import io
import json
import urllib.request

from openpyxl import load_workbook

BASE = Path(__file__).parent
SOURCES = BASE / "sources"
SOURCES.mkdir(parents=True, exist_ok=True)
URLS = {
    "native_campaign_template.csv": "https://www.gstatic.com/adwords/campaignmgmt/templates/campaign_template.csv",
    "native_ad_group_template.csv": "https://www.gstatic.com/adwords/campaignmgmt/templates/ad_group_template.csv",
    "native_responsive_search_ad_template.csv": "https://www.gstatic.com/adwords/campaignmgmt/templates/responsive_search_ad_template.csv",
    "native_keyword_template.csv": "https://www.gstatic.com/adwords/campaignmgmt/templates/keyword_template.csv",
    "native_ad_group_negative_keyword_template.csv": "https://www.gstatic.com/adwords/campaignmgmt/templates/ad_group_negative_keyword_template.csv",
    "help_create_campaign.xlsx": "https://storage.googleapis.com/support-kms-prod/lx4kKazVwHuqcFVQPZAKJJ5EUQAV5VBBMrn2",
    "help_create_sitelink.csv": "https://storage.googleapis.com/support-kms-prod/er1LbgoAEsIctW1koPH6f4fDr7NNVDCKWPvh",
    "help_create_callout.csv": "https://storage.googleapis.com/support-kms-prod/Tqa4nGJSfFLXLSuw0Z4G17HVZJ1nAQMpYH37",
    "help_create_snippet.csv": "https://storage.googleapis.com/support-kms-prod/EotlZJ2OljTMZYsB4Qx0a14sIJYwvqTU8Gpu",
    "help_create_keywords.xlsx": "https://storage.googleapis.com/support-kms-prod/ZPXjGm9HgYzegUH8aExrVUm3tHTp5rrvAc9a",
    "help_create_ad_groups.xlsx": "https://storage.googleapis.com/support-kms-prod/eSiqsHfH2KNrnD4wnYuMX49DlTkbbz2C10Bs",
    "help_create_rsa.xlsx": "https://storage.googleapis.com/support-kms-prod/4AQnEPOMnm95I6VILqWdg1xypQ9KlaXG2IgY",
}

def fetch(item):
    name, url = item
    path = SOURCES / name
    if path.exists():
        data = path.read_bytes()
    else:
        request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(request, timeout=30) as response:
            data = response.read()
        path.write_bytes(data)
    if name.endswith(".xlsx"):
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        tables = {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)][:12] for sheet in book.worksheets}
    else:
        tables = {"CSV": list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))[:14]}
    return {"file": name, "url": url, "sha256": hashlib.sha256(data).hexdigest(), "tables": tables}

with concurrent.futures.ThreadPoolExecutor(max_workers=6) as executor:
    results = list(executor.map(fetch, URLS.items()))
(SOURCES / "template_manifest.json").write_text(json.dumps(results, ensure_ascii=False, indent=2, default=str) + "\n")
for result in results:
    print(json.dumps(result, ensure_ascii=False, default=str))
